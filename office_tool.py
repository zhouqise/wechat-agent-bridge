"""Validate and extract Office content. Never executes macros or external links."""
import json, sys, zipfile, pathlib

def inspect(filename):
    p=pathlib.Path(filename); ext=p.suffix.lower()
    if ext not in ('.xlsx','.docx','.pptx'): raise ValueError('format')
    if p.stat().st_size>20*1024*1024: raise ValueError('size')
    with zipfile.ZipFile(p) as z:
        entries=z.infolist()
        if len(entries)>5000 or sum(i.file_size for i in entries)>80*1024*1024: raise ValueError('size')
        if len({i.filename for i in entries})!=len(entries): raise ValueError('duplicate')
        for i in entries:
            if i.file_size>16*1024*1024 or i.flag_bits & 1: raise ValueError('size/encryption')
            if 'vbaproject' in i.filename.lower(): raise ValueError('macro')
            if i.filename.endswith(('.xml','.rels')):
                content=z.read(i)
                if b'<!DOCTYPE' in content or b'<!ENTITY' in content: raise ValueError('entity')
        required={'.xlsx':'xl/workbook.xml','.docx':'word/document.xml','.pptx':'ppt/presentation.xml'}[ext]
        if required not in z.namelist(): raise ValueError('format')
    lines=[]; truncated=False
    def add(text):
        nonlocal truncated
        if len(lines)>=3000 or sum(map(len,lines))>30000: truncated=True;return False
        lines.append(str(text)[:10000]);return True
    if ext=='.xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(p,read_only=True,data_only=False,keep_links=False)
        for ws in wb.worksheets[:30]:
            add('工作表：'+ws.title)
            for row in ws.iter_rows(max_row=min(ws.max_row or 1,2000),max_col=min(ws.max_column or 1,100)):
                values=['%s=%s'%(c.coordinate,c.value) for c in row if c.value is not None]
                if values and not add(' | '.join(values)): break
            if (ws.max_row or 0)>2000 or (ws.max_column or 0)>100: truncated=True
        if len(wb.worksheets)>30: truncated=True
        wb.close()
    elif ext=='.docx':
        from docx import Document
        doc=Document(p)
        for paragraph in doc.paragraphs:
            if not add(paragraph.text): break
        for i,table in enumerate(doc.tables):
            if not add('表格 %d'%(i+1)): break
            for row in table.rows:
                if not add(' | '.join(c.text for c in row.cells)): break
    else:
        from pptx import Presentation
        deck=Presentation(p)
        for i,slide in enumerate(deck.slides):
            if not add('幻灯片 %d'%(i+1)): break
            for shape in slide.shapes:
                if shape.has_text_frame: add(shape.text)
                if shape.has_table:
                    for row in shape.table.rows: add(' | '.join(c.text for c in row.cells))
            if slide.has_notes_slide: add('备注：'+slide.notes_slide.notes_text_frame.text)
    return {'format':ext,'text':'\n'.join(lines)[:30000],'truncated':truncated,'notice':'提取文本、表格和公式；未渲染版式、图表或图片。Excel 公式不在此重新计算。'}

if __name__=='__main__':
    try: print(json.dumps(inspect(sys.argv[1]),ensure_ascii=True))
    except Exception: print(json.dumps({'error':'文档无法解析：请确认格式正确、未加密且未超出大小或解压限制'},ensure_ascii=True));sys.exit(1)
